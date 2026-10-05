from __future__ import annotations

import json
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

from capabilities.OllamaReceiptAcquisitionPipeline.ollama_client import (
    OllamaReceiptAcquisitionError,
    run_local_json_prompt,
)
from utils.constants import DATA_DIR
from utils.DataBaseBuilder.excel.category_manager import (
    CATEGORY_HEADER,
    CATEGORY_MANAGER_SHEET,
    PRODUCT_HEADER_PREFIX,
    SUB_CATEGORY_HEADER,
    _is_na,
    _normalized,
    _text,
    _verified_atomic_save,
)
from utils.DataBaseBuilder.excel.purchase_history import (
    CATEGORY_COLUMN,
    COMMON_NAME_COLUMN,
    PRODUCT_COLUMN,
    PURCHASE_SHEET,
    SKU_COLUMN,
    STORE_COLUMN,
    WORKBOOK_PATH,
    _ensure_purchase_schema,
)
from utils.DataBaseBuilder.purchase_record import NA

PROMPT_PATH = DATA_DIR / "prompts" / "dev_prompts" / "shopgraph_ollama_category_completion_prompt.txt"
BATCH_SIZE = 15
MAX_TAXONOMY_REFERENCES = 250
MAX_PURCHASE_EXAMPLES = 8
MAX_CATEGORY_LENGTH = 80


def _load_workbook_for_completion():
    path = Path(WORKBOOK_PATH).expanduser().resolve()
    if not path.exists():
        raise FileNotFoundError(f"Purchase History workbook was not found:\n{path}")
    workbook = load_workbook(path)
    required = {PURCHASE_SHEET, CATEGORY_MANAGER_SHEET}
    missing = required - set(workbook.sheetnames)
    if missing:
        workbook.close()
        raise ValueError("Workbook is missing required sheet(s): " + ", ".join(sorted(missing)))
    purchase = workbook[PURCHASE_SHEET]
    _ensure_purchase_schema(purchase)
    manager = workbook[CATEGORY_MANAGER_SHEET]
    if _text(manager.cell(1, 1).value) != CATEGORY_HEADER or _text(manager.cell(1, 2).value) != SUB_CATEGORY_HEADER:
        workbook.close()
        raise ValueError('Category Manager must begin with "Category" and "Sub-Category".')
    return path, workbook, purchase, manager


def _authoritative_taxonomy(manager) -> tuple[list[dict], dict[str, set[str]], dict[str, str]]:
    refs = []
    by_subcategory: dict[str, set[str]] = defaultdict(set)
    category_spelling: dict[str, str] = {}
    for row in range(2, manager.max_row + 1):
        category = _text(manager.cell(row, 1).value)
        subcategory = _text(manager.cell(row, 2).value)
        if not subcategory or not category or _is_na(category):
            continue
        by_subcategory[_normalized(subcategory)].add(category)
        category_spelling.setdefault(_normalized(category), category)
        refs.append({"sub_category": subcategory, "category": category})
    return refs[:MAX_TAXONOMY_REFERENCES], by_subcategory, category_spelling


def _products_for_manager_row(manager, row: int) -> list[str]:
    products = []
    for column in range(3, manager.max_column + 1):
        header = _text(manager.cell(1, column).value)
        if header and not header.startswith(PRODUCT_HEADER_PREFIX):
            continue
        value = _text(manager.cell(row, column).value)
        if value and not _is_na(value):
            products.append(value)
    return products


def _purchase_examples(purchase, subcategory: str) -> list[dict]:
    examples = []
    seen = set()
    target = _normalized(subcategory)
    for row in range(2, purchase.max_row + 1):
        if _normalized(purchase.cell(row, CATEGORY_COLUMN).value) != target:
            continue
        example = {
            "store": _text(purchase.cell(row, STORE_COLUMN).value),
            "six_digit_sku": _text(purchase.cell(row, SKU_COLUMN).value),
            "product": _text(purchase.cell(row, PRODUCT_COLUMN).value),
            "common_name": _text(purchase.cell(row, COMMON_NAME_COLUMN).value),
            "sub_category": _text(purchase.cell(row, CATEGORY_COLUMN).value),
        }
        key = tuple(example.values())
        if key in seen:
            continue
        seen.add(key)
        examples.append(example)
        if len(examples) >= MAX_PURCHASE_EXAMPLES:
            break
    return examples


def _target_payload(manager, purchase, row: int) -> dict:
    subcategory = _text(manager.cell(row, 2).value)
    return {
        "excel_row": row,
        "sub_category": subcategory,
        "products": _products_for_manager_row(manager, row),
        "purchase_history_examples": _purchase_examples(purchase, subcategory),
    }


def _build_prompt(rows: list[dict], taxonomy: list[dict]) -> str:
    if not PROMPT_PATH.exists():
        raise FileNotFoundError(f"Ollama Category completion prompt was not found:\n{PROMPT_PATH.resolve()}")
    base = PROMPT_PATH.read_text(encoding="utf-8").strip()
    payload = {"existing_taxonomy": taxonomy, "rows_to_complete": rows}
    return base + "\n\nINPUT JSON:\n" + json.dumps(payload, ensure_ascii=False, indent=2)


def _clean_category(value: object, row: int, category_spelling: dict[str, str]) -> str:
    if not isinstance(value, str):
        raise ValueError(f"Ollama returned a non-string Category for row {row}.")
    category = value.strip()
    if not category or _is_na(category):
        raise ValueError(f"Ollama returned blank/NA Category for row {row}.")
    if "\n" in category or "\r" in category or "\t" in category:
        raise ValueError(f"Ollama returned malformed Category for row {row}.")
    if len(category) > MAX_CATEGORY_LENGTH:
        raise ValueError(f"Ollama returned an excessively long Category for row {row}.")
    return category_spelling.get(_normalized(category), category)


def _validate_response(response: dict, requested: list[dict], category_spelling: dict[str, str]) -> dict[int, str]:
    rows = response.get("rows")
    if not isinstance(rows, list):
        raise ValueError('Ollama response must contain a "rows" list.')
    expected = {item["excel_row"]: item["sub_category"] for item in requested}
    results: dict[int, str] = {}
    for item in rows:
        if not isinstance(item, dict):
            raise ValueError("Every Ollama rows entry must be an object.")
        row = item.get("excel_row")
        subcategory = item.get("sub_category")
        if not isinstance(row, int) or row not in expected:
            raise ValueError(f"Ollama returned an unexpected Excel row: {row!r}")
        if row in results:
            raise ValueError(f"Ollama returned duplicate Excel row {row}.")
        if not isinstance(subcategory, str) or _normalized(subcategory) != _normalized(expected[row]):
            raise ValueError(f"Ollama changed Sub-Category identity for row {row}.")
        results[row] = _clean_category(item.get("category"), row, category_spelling)
    missing = sorted(set(expected) - set(results))
    if missing:
        raise ValueError(f"Ollama omitted requested Excel rows: {missing}")
    return results


def _remaining_missing(manager) -> list[int]:
    result = []
    for row in range(2, manager.max_row + 1):
        subcategory = _text(manager.cell(row, 2).value)
        category = _text(manager.cell(row, 1).value)
        if subcategory and (not category or _is_na(category)):
            result.append(row)
    return result


def run_ollama_category_manager_completion() -> dict:
    print("\n=== Ollama Category Manager Completion ===\n")
    print("[1/5] Load Category Manager")
    workbook_path, workbook, purchase, manager = _load_workbook_for_completion()
    try:
        print("[2/5] Identify Missing Categories")
        targets = _remaining_missing(manager)
        if not targets:
            print("\n[OK] Category Manager contains no blank/NA Categories.")
            print("No changes were required.")
            return {"updated_rows": 0, "remaining_rows": []}

        taxonomy, by_subcategory, category_spelling = _authoritative_taxonomy(manager)
        original_subcategories = {row: _text(manager.cell(row, 2).value) for row in targets}
        proposed: dict[int, str] = {}
        deterministic: set[int] = set()
        ai_rows = []
        for row in targets:
            payload = _target_payload(manager, purchase, row)
            matches = by_subcategory.get(_normalized(payload["sub_category"]), set())
            if len(matches) == 1:
                proposed[row] = next(iter(matches))
                deterministic.add(row)
            else:
                ai_rows.append(payload)

        print(f"\nFound {len(targets)} Category Manager rows with blank/NA Category.")
        if deterministic:
            print(f"[INFO] Reused {len(deterministic)} exact existing Sub-Category -> Category mappings.")

        print("\n[3/5] Analyze Missing Categories with Ollama")
        for start in range(0, len(ai_rows), BATCH_SIZE):
            batch = ai_rows[start:start + BATCH_SIZE]
            response, _metadata = run_local_json_prompt(_build_prompt(batch, taxonomy))
            proposed.update(_validate_response(response, batch, category_spelling))

        print("\nProposed Category assignments:\n")
        for row in targets:
            source = "deterministic" if row in deterministic else "Ollama"
            print(f"Row {row} | {original_subcategories[row]} -> {proposed[row]} [{source}]")

        print("\n[4/5] Validate Ollama Results")
        if set(proposed) != set(targets):
            missing = sorted(set(targets) - set(proposed))
            raise ValueError(f"Completion did not produce every target row: {missing}")
        for row in targets:
            current_category = _text(manager.cell(row, 1).value)
            current_subcategory = _text(manager.cell(row, 2).value)
            if current_category and not _is_na(current_category):
                raise ValueError(f"Refusing to overwrite non-NA Category Manager row {row}.")
            if _normalized(current_subcategory) != _normalized(original_subcategories[row]):
                raise ValueError(f"Sub-Category identity changed before write for row {row}.")
            proposed[row] = _clean_category(proposed[row], row, category_spelling)

        print(f"\n[OK] {len(proposed)}/{len(targets)} Category assignments validated.")
        print("\n[5/5] Update Category Manager")
        for row, category in proposed.items():
            manager.cell(row, 1, value=category)
        _verified_atomic_save(workbook, workbook_path)
        print("\n[OK] Category Manager updated.")
        print(f"Rows updated: {len(proposed)}")
        print(f"Workbook:\n{workbook_path}")
    finally:
        workbook.close()

    verify = load_workbook(workbook_path, read_only=True)
    try:
        remaining = _remaining_missing(verify[CATEGORY_MANAGER_SHEET])
    finally:
        verify.close()
    if remaining:
        print("\n[WARNING] Category Manager still contains Sub-Categories without a valid Category.")
        print("Rows: " + ", ".join(str(row) for row in remaining))
    else:
        print("\n[OK] Category Manager Category validation passed.")
        print("Category Manager rows with blank/NA Category: 0")
    return {"updated_rows": len(proposed), "remaining_rows": remaining}


def main() -> None:
    try:
        run_ollama_category_manager_completion()
    except (FileNotFoundError, PermissionError, OSError, ValueError, OllamaReceiptAcquisitionError) as error:
        print(f"\n[ERROR] Ollama Category Manager completion failed:\n{error}")


if __name__ == "__main__":
    main()

