from __future__ import annotations

import json
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

from capabilities.OllamaReceiptAcquisitionPipeline.ollama_client import OllamaReceiptAcquisitionError, run_local_json_prompt
from utils.constants import DATA_DIR
from utils.DataBaseBuilder.excel.category_manager import CATEGORY_MANAGER_SHEET, _is_na, _normalized, _text, _verified_atomic_save
from utils.DataBaseBuilder.excel.purchase_history import CATEGORY_COLUMN, COMMON_NAME_COLUMN, FIXED_HEADERS, PRODUCT_COLUMN, PURCHASE_SHEET, STORE_COLUMN, WORKBOOK_PATH, _ensure_purchase_schema

PROMPT_PATH = DATA_DIR / "prompts" / "dev_prompts" / "shopgraph_ollama_purchase_history_subcategory_completion_prompt.txt"
BATCH_SIZE = 20
MAX_SUBCATEGORY_LENGTH = 100


def _load():
    path = Path(WORKBOOK_PATH).expanduser().resolve()
    if not path.exists(): raise FileNotFoundError(f"Purchase History workbook was not found:\n{path}")
    wb = load_workbook(path)
    if PURCHASE_SHEET not in wb.sheetnames:
        wb.close(); raise ValueError(f'Workbook is missing required sheet: "{PURCHASE_SHEET}".')
    sh = wb[PURCHASE_SHEET]; _ensure_purchase_schema(sh)
    actual=[_text(sh.cell(1,c).value) for c in range(1,len(FIXED_HEADERS)+1)]
    if actual != FIXED_HEADERS:
        wb.close(); raise ValueError("Purchase History has an unrecognized column layout.")
    manager = wb[CATEGORY_MANAGER_SHEET] if CATEGORY_MANAGER_SHEET in wb.sheetnames else None
    return path, wb, sh, manager


def _targets(sh):
    rows=[]
    for r in range(2, sh.max_row+1):
        product=_text(sh.cell(r,PRODUCT_COLUMN).value)
        if not product: continue
        value=_text(sh.cell(r,CATEGORY_COLUMN).value)
        if not value or _is_na(value): rows.append(r)
    return rows


def _existing(sh):
    by_common=defaultdict(set); by_product=defaultdict(set); spelling={}
    for r in range(2,sh.max_row+1):
        sub=_text(sh.cell(r,CATEGORY_COLUMN).value)
        if not sub or _is_na(sub): continue
        spelling.setdefault(_normalized(sub),sub)
        common=_text(sh.cell(r,COMMON_NAME_COLUMN).value); product=_text(sh.cell(r,PRODUCT_COLUMN).value)
        if common and not _is_na(common): by_common[_normalized(common)].add(sub)
        if product: by_product[_normalized(product)].add(sub)
    return by_common,by_product,spelling


def _taxonomy(manager, limit=250):
    if manager is None: return []
    refs=[]
    for r in range(2,manager.max_row+1):
        cat=_text(manager.cell(r,1).value); sub=_text(manager.cell(r,2).value)
        if sub and not _is_na(sub): refs.append({"category":cat,"sub_category":sub})
        if len(refs)>=limit: break
    return refs


def _payload(sh,r):
    return {"excel_row":r,"store":_text(sh.cell(r,STORE_COLUMN).value),"product":_text(sh.cell(r,PRODUCT_COLUMN).value),"common_name":_text(sh.cell(r,COMMON_NAME_COLUMN).value)}


def _build(rows,taxonomy):
    if not PROMPT_PATH.exists(): raise FileNotFoundError(f"Prompt file was not found:\n{PROMPT_PATH.resolve()}")
    return PROMPT_PATH.read_text(encoding='utf-8').strip()+"\n\nINPUT JSON:\n"+json.dumps({"existing_taxonomy":taxonomy,"rows_to_complete":rows},ensure_ascii=False,indent=2)


def _clean(v,row,spelling):
    if not isinstance(v,str): raise ValueError(f"Ollama returned a non-string Sub-Category for row {row}.")
    v=v.strip()
    if not v or _is_na(v): raise ValueError(f"Ollama returned blank/NA Sub-Category for row {row}.")
    if any(x in v for x in '\n\r\t') or len(v)>MAX_SUBCATEGORY_LENGTH: raise ValueError(f"Ollama returned malformed Sub-Category for row {row}.")
    return spelling.get(_normalized(v),v)


def _validate(resp,requested,spelling):
    rows=resp.get('rows'); expected={x['excel_row'] for x in requested}
    if not isinstance(rows,list): raise ValueError('Ollama response must contain a "rows" list.')
    result={}
    for item in rows:
        if not isinstance(item,dict): raise ValueError('Every Ollama rows entry must be an object.')
        r=item.get('excel_row')
        if not isinstance(r,int) or r not in expected: raise ValueError(f"Ollama returned an unexpected Excel row: {r!r}")
        if r in result: raise ValueError(f"Ollama returned duplicate Excel row {r}.")
        result[r]=_clean(item.get('sub_category'),r,spelling)
    missing=sorted(expected-set(result))
    if missing: raise ValueError(f"Ollama omitted requested Excel rows: {missing}")
    return result


def run_ollama_purchase_history_subcategory_completion() -> dict:
    print("\n=== Ollama Purchase History Sub-Category Blank/NA Completion ===\n")
    path,wb,sh,manager=_load()
    try:
        targets=_targets(sh)
        if not targets:
            print("[OK] Purchase History contains no populated rows with blank/NA Sub-Category."); return {"updated_rows":0,"remaining_rows":[]}
        by_common,by_product,spelling=_existing(sh); taxonomy=_taxonomy(manager)
        original={r:(_text(sh.cell(r,PRODUCT_COLUMN).value),_text(sh.cell(r,COMMON_NAME_COLUMN).value)) for r in targets}
        proposed={}; deterministic=set(); ai=[]
        for r in targets:
            payload=_payload(sh,r); common=payload['common_name']; product=payload['product']; matches=set()
            if common and not _is_na(common): matches |= by_common.get(_normalized(common),set())
            pm=by_product.get(_normalized(product),set())
            if len(matches)==1: pass
            elif not matches and len(pm)==1: matches=pm
            else: matches=set()
            if len(matches)==1: proposed[r]=next(iter(matches)); deterministic.add(r)
            else: ai.append(payload)
        print(f"Found {len(targets)} populated rows with blank/NA Sub-Category.")
        print(f"[INFO] Deterministic assignments: {len(deterministic)}")
        print(f"[INFO] Rows requiring Ollama: {len(ai)}")
        for start in range(0,len(ai),BATCH_SIZE):
            batch=ai[start:start+BATCH_SIZE]; resp,_=run_local_json_prompt(_build(batch,taxonomy)); proposed.update(_validate(resp,batch,spelling))
        if set(proposed)!=set(targets): raise ValueError(f"Completion did not produce every target row: {sorted(set(targets)-set(proposed))}")
        for r in targets:
            if (_text(sh.cell(r,PRODUCT_COLUMN).value),_text(sh.cell(r,COMMON_NAME_COLUMN).value)) != original[r]: raise ValueError(f"Purchase History row identity changed before write for row {r}.")
            current=_text(sh.cell(r,CATEGORY_COLUMN).value)
            if current and not _is_na(current): raise ValueError(f"Refusing to overwrite non-NA Sub-Category at Purchase History row {r}.")
            proposed[r]=_clean(proposed[r],r,spelling)
        for r,v in proposed.items(): sh.cell(r,CATEGORY_COLUMN,value=v)
        _verified_atomic_save(wb,path)
        print(f"\n[OK] Purchase History updated. Rows updated: {len(proposed)}\nWorkbook:\n{path}")
    finally: wb.close()
    verify=load_workbook(path,read_only=True)
    try: remaining=_targets(verify[PURCHASE_SHEET])
    finally: verify.close()
    if remaining: print("\n[WARNING] Purchase History still contains blank/NA Sub-Category rows: "+", ".join(map(str,remaining)))
    else: print("\n[OK] Purchase History Sub-Category validation passed.\nPopulated rows with blank/NA Sub-Category: 0")
    return {"updated_rows":len(proposed),"remaining_rows":remaining}


def main():
    try: run_ollama_purchase_history_subcategory_completion()
    except (FileNotFoundError,PermissionError,OSError,ValueError,OllamaReceiptAcquisitionError) as e: print(f"\n[ERROR] Ollama Purchase History Sub-Category completion failed:\n{e}")

if __name__=='__main__': main()
