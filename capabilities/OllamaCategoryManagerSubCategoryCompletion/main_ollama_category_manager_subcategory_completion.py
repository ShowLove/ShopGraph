from __future__ import annotations
import json
from collections import defaultdict
from pathlib import Path
from openpyxl import load_workbook
from capabilities.OllamaReceiptAcquisitionPipeline.ollama_client import OllamaReceiptAcquisitionError, run_local_json_prompt
from utils.constants import DATA_DIR
from utils.DataBaseBuilder.excel.category_manager import CATEGORY_HEADER,CATEGORY_MANAGER_SHEET,PRODUCT_HEADER_PREFIX,SUB_CATEGORY_HEADER,_is_na,_normalized,_text,_verified_atomic_save
from utils.DataBaseBuilder.excel.purchase_history import CATEGORY_COLUMN,COMMON_NAME_COLUMN,PRODUCT_COLUMN,PURCHASE_SHEET,WORKBOOK_PATH,_ensure_purchase_schema

PROMPT_PATH=DATA_DIR/'prompts'/'dev_prompts'/'shopgraph_ollama_category_manager_subcategory_completion_prompt.txt'
BATCH_SIZE=15; MAX_SUBCATEGORY_LENGTH=100

def _load():
    path=Path(WORKBOOK_PATH).expanduser().resolve()
    if not path.exists(): raise FileNotFoundError(f"Purchase History workbook was not found:\n{path}")
    wb=load_workbook(path); required={PURCHASE_SHEET,CATEGORY_MANAGER_SHEET}; missing=required-set(wb.sheetnames)
    if missing: wb.close(); raise ValueError('Workbook is missing required sheet(s): '+', '.join(sorted(missing)))
    purchase=wb[PURCHASE_SHEET]; _ensure_purchase_schema(purchase); manager=wb[CATEGORY_MANAGER_SHEET]
    if _text(manager.cell(1,1).value)!=CATEGORY_HEADER or _text(manager.cell(1,2).value)!=SUB_CATEGORY_HEADER: wb.close(); raise ValueError('Category Manager must begin with "Category" and "Sub-Category".')
    return path,wb,purchase,manager

def _products(manager,row):
    vals=[]
    for c in range(3,manager.max_column+1):
        h=_text(manager.cell(1,c).value)
        if h and not h.startswith(PRODUCT_HEADER_PREFIX): continue
        v=_text(manager.cell(row,c).value)
        if v and not _is_na(v): vals.append(v)
    return vals

def _targets(manager):
    result=[]
    for r in range(2,manager.max_row+1):
        sub=_text(manager.cell(r,2).value)
        if sub and not _is_na(sub): continue
        cat=_text(manager.cell(r,1).value); products=_products(manager,r)
        if (cat and not _is_na(cat)) or products: result.append(r)
    return result

def _purchase_maps(purchase):
    by_common=defaultdict(set); by_product=defaultdict(set); spelling={}
    for r in range(2,purchase.max_row+1):
        sub=_text(purchase.cell(r,CATEGORY_COLUMN).value)
        if not sub or _is_na(sub): continue
        spelling.setdefault(_normalized(sub),sub)
        common=_text(purchase.cell(r,COMMON_NAME_COLUMN).value); product=_text(purchase.cell(r,PRODUCT_COLUMN).value)
        if common and not _is_na(common): by_common[_normalized(common)].add(sub)
        if product: by_product[_normalized(product)].add(sub)
    return by_common,by_product,spelling

def _taxonomy(manager,limit=250):
    refs=[]
    for r in range(2,manager.max_row+1):
        cat=_text(manager.cell(r,1).value); sub=_text(manager.cell(r,2).value)
        if sub and not _is_na(sub): refs.append({'category':cat,'sub_category':sub})
        if len(refs)>=limit: break
    return refs

def _build(rows,taxonomy):
    if not PROMPT_PATH.exists(): raise FileNotFoundError(f"Prompt file was not found:\n{PROMPT_PATH.resolve()}")
    return PROMPT_PATH.read_text(encoding='utf-8').strip()+"\n\nINPUT JSON:\n"+json.dumps({'existing_taxonomy':taxonomy,'rows_to_complete':rows},ensure_ascii=False,indent=2)

def _clean(v,row,spelling):
    if not isinstance(v,str): raise ValueError(f"Ollama returned a non-string Sub-Category for row {row}.")
    v=v.strip()
    if not v or _is_na(v): raise ValueError(f"Ollama returned blank/NA Sub-Category for row {row}.")
    if any(x in v for x in '\n\r\t') or len(v)>MAX_SUBCATEGORY_LENGTH: raise ValueError(f"Ollama returned malformed Sub-Category for row {row}.")
    return spelling.get(_normalized(v),v)

def _validate(resp,requested,spelling):
    rows=resp.get('rows'); expected={x['excel_row'] for x in requested}
    if not isinstance(rows,list): raise ValueError('Ollama response must contain a "rows" list.')
    result={}
    for item in rows:
        if not isinstance(item,dict): raise ValueError('Every Ollama rows entry must be an object.')
        r=item.get('excel_row')
        if not isinstance(r,int) or r not in expected: raise ValueError(f"Ollama returned an unexpected Excel row: {r!r}")
        if r in result: raise ValueError(f"Ollama returned duplicate Excel row {r}.")
        result[r]=_clean(item.get('sub_category'),r,spelling)
    missing=sorted(expected-set(result))
    if missing: raise ValueError(f"Ollama omitted requested Excel rows: {missing}")
    return result

def run_ollama_category_manager_subcategory_completion()->dict:
    print("\n=== Ollama Category Manager Sub-Category Blank/NA Completion ===\n")
    path,wb,purchase,manager=_load()
    try:
        targets=_targets(manager)
        if not targets: print('[OK] Category Manager contains no meaningful rows with blank/NA Sub-Category.'); return {'updated_rows':0,'remaining_rows':[]}
        by_common,by_product,spelling=_purchase_maps(purchase); taxonomy=_taxonomy(manager); original={r:(_text(manager.cell(r,1).value),tuple(_products(manager,r))) for r in targets}
        proposed={}; deterministic=set(); ai=[]
        for r in targets:
            products=_products(manager,r); matches=set(); ambiguous=False
            for product in products:
                m=by_common.get(_normalized(product),set()) or by_product.get(_normalized(product),set())
                if len(m)>1: ambiguous=True; break
                if len(m)==1: matches |= m
            if not ambiguous and len(matches)==1: proposed[r]=next(iter(matches)); deterministic.add(r)
            else: ai.append({'excel_row':r,'category':_text(manager.cell(r,1).value),'products':products})
        print(f"Found {len(targets)} Category Manager rows with blank/NA Sub-Category.")
        print(f"[INFO] Deterministic assignments: {len(deterministic)}")
        print(f"[INFO] Rows requiring Ollama: {len(ai)}")
        for start in range(0,len(ai),BATCH_SIZE):
            batch=ai[start:start+BATCH_SIZE]; resp,_=run_local_json_prompt(_build(batch,taxonomy)); proposed.update(_validate(resp,batch,spelling))
        if set(proposed)!=set(targets): raise ValueError(f"Completion did not produce every target row: {sorted(set(targets)-set(proposed))}")
        for r in targets:
            if (_text(manager.cell(r,1).value),tuple(_products(manager,r))) != original[r]: raise ValueError(f"Category Manager row identity changed before write for row {r}.")
            cur=_text(manager.cell(r,2).value)
            if cur and not _is_na(cur): raise ValueError(f"Refusing to overwrite non-NA Category Manager Sub-Category row {r}.")
            proposed[r]=_clean(proposed[r],r,spelling)
        for r,v in proposed.items(): manager.cell(r,2,value=v)
        _verified_atomic_save(wb,path)
        print(f"\n[OK] Category Manager Sub-Categories updated. Rows updated: {len(proposed)}\nWorkbook:\n{path}")
    finally: wb.close()
    verify=load_workbook(path,read_only=True)
    try: remaining=_targets(verify[CATEGORY_MANAGER_SHEET])
    finally: verify.close()
    if remaining: print('\n[WARNING] Category Manager still contains blank/NA Sub-Category rows: '+', '.join(map(str,remaining)))
    else: print('\n[OK] Category Manager Sub-Category validation passed.\nCategory Manager rows with blank/NA Sub-Category: 0')
    return {'updated_rows':len(proposed),'remaining_rows':remaining}

def main():
    try: run_ollama_category_manager_subcategory_completion()
    except (FileNotFoundError,PermissionError,OSError,ValueError,OllamaReceiptAcquisitionError) as e: print(f"\n[ERROR] Ollama Category Manager Sub-Category completion failed:\n{e}")
if __name__=='__main__': main()
